"""CST workbook format with serialized, atomic file access on I/O workers."""
from pathlib import Path
from threading import Lock
import os
import shutil
import tempfile
from openpyxl import load_workbook
from python.app.history.snapshot_save_qt import append_snapshot_xlsx

_WORKBOOK_LOCK=Lock()


def save_snapshot(path,signals,note,timestamp):
    target=Path(path).resolve()
    target.parent.mkdir(parents=True,exist_ok=True)
    with _WORKBOOK_LOCK:
        with tempfile.TemporaryDirectory(prefix="snapshot-",dir=target.parent) as temporary:
            staged=Path(temporary)/target.name
            if target.exists():shutil.copyfile(target,staged)
            append_snapshot_xlsx(str(staged),signals,note,ts_epoch=timestamp,scaled_only=False)
            os.replace(staged,target)
    return str(target)


def read_snapshots(path):
    with _WORKBOOK_LOCK:
        workbook=load_workbook(path,data_only=True,read_only=True)
        try:
            sheet=workbook.active
            headers=[str(value).strip() if value is not None else "" for value in next(sheet.iter_rows(min_row=1,max_row=1,values_only=True))]
            rows=[dict(zip(headers,row)) for row in sheet.iter_rows(min_row=2,values_only=True) if row and any(value is not None for value in row)]
        finally:
            workbook.close()
    return headers,rows
